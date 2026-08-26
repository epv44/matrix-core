import io
import logging
import os
import stat
from datetime import datetime, timezone

import openpyxl
from openpyxl.utils import get_column_letter

from finance.categorization.processor import process_monthly_data
from finance.services.categorization import resolve_category_for_transaction
from finance.storage.repository import (
    get_category_rules,
    get_sheet_layout,
    get_taxonomy_mappings,
    get_transactions_by_date_range,
)

logger = logging.getLogger(__name__)
EXPORT_FILENAME_TEMPLATE = "finance_report_{year}_{month:02d}_{day:02d}.xlsx"


def _col_for_category(i: int) -> int:
    return 1 + i * 3


def _remove_if_not_none(wb):
    ws = wb.active
    if ws is not None:
        wb.remove(ws)


def _write_month_sheet(wb, year, month, month_data):
    ws = wb.create_sheet(title=f"{year}-{month:02d}")
    summary = month_data["summary"]
    ws.cell(row=1, column=1, value="Total Spent:")
    ws.cell(row=1, column=2, value=summary["total_spent"])
    ws.cell(row=1, column=3, value="Total Made:")
    ws.cell(row=1, column=4, value=summary["total_made"])
    ws.cell(row=1, column=5, value="Net P/L:")
    ws.cell(row=1, column=6, value=summary["net"])
    ws.cell(row=1, column=7, value="% Saved:")
    ws.cell(row=1, column=8, value=summary["saved_pct"])
    cols = list(month_data["expenses"].keys())
    header_row = 4
    for idx, name in enumerate(cols):
        ws.cell(row=header_row, column=_col_for_category(idx), value=name)
    expenses = month_data["expenses"]
    for sc in cols:
        expenses.setdefault(sc, [])
    data_start = 5
    max_txns = max((len(v) for v in expenses.values()), default=0)
    for off in range(max_txns):
        row = data_start + off
        for c_idx, c_name in enumerate(cols):
            col = _col_for_category(c_idx)
            lst = expenses.get(c_name, [])
            if off < len(lst):
                name, amt = lst[off]
                ws.cell(row=row, column=col, value=name)
                ws.cell(row=row, column=col + 1, value=float(amt))
    total_row = data_start + max_txns
    for c_idx, c_name in enumerate(cols):
        col = _col_for_category(c_idx)
        ws.cell(row=total_row, column=col, value="Total:")
        lst = expenses.get(c_name, [])
        if lst:
            first, last = data_start, data_start + len(lst) - 1
            letter = get_column_letter(col + 1)
            ws.cell(
                row=total_row,
                column=col + 1,
                value=f"=SUM({letter}{first}:{letter}{last})",
            )


def _exporter_config(conn):

    layout = get_sheet_layout(conn)
    rules = get_category_rules(conn)
    mappings = get_taxonomy_mappings(conn)
    return (
        layout,
        lambda t: (
            resolve_category_for_transaction(conn, t, rules=rules, mappings=mappings)
            or "MISC."
        ),
    )


def export_to_excel(conn, output_dir, year_start, month_start, year_end, month_end):
    now = datetime.now(timezone.utc)
    filename = EXPORT_FILENAME_TEMPLATE.format(
        year=now.year, month=now.month, day=now.day
    )
    output_path = os.path.join(output_dir, filename)
    layout, category_resolver = _exporter_config(conn)
    raw = {}
    y, m = year_start, month_start
    while (y < year_end) or (y == year_end and m <= month_end):
        txns = _month_transactions(conn, y, m)
        if txns:
            raw[(y, m)] = txns
        m += 1
        if m > 12:
            m = 1
            y += 1
    processed = process_monthly_data(
        raw, category_resolver=category_resolver, sheet_categories=layout
    )
    wb = openpyxl.Workbook()
    _remove_if_not_none(wb)
    for year, month in sorted(processed.keys()):
        _write_month_sheet(wb, year, month, processed[(year, month)])
    wb.save(output_path)
    os.chmod(output_path, stat.S_IRUSR | stat.S_IWUSR)
    logger.info("Exported to %s", output_path)
    return output_path


def export_from_rows(monthly_txns, output_path):
    processed = process_monthly_data(monthly_txns)
    wb = openpyxl.Workbook()
    _remove_if_not_none(wb)
    for year, month in sorted(processed.keys()):
        _write_month_sheet(wb, year, month, processed[(year, month)])
    wb.save(output_path)
    os.chmod(output_path, stat.S_IRUSR | stat.S_IWUSR)
    logger.info("Exported to %s", output_path)
    return output_path


def export_months_to_bytes(conn, months_data):
    # months_data is processed dict {(y,m): {summary,expenses,...}} OR raw dict
    wb = openpyxl.Workbook()
    _remove_if_not_none(wb)
    # normalize: if values contain transactions list, process
    to_write = {}
    has_summary = any(
        isinstance(v, dict) and "summary" in v for v in months_data.values()
    )
    if has_summary:
        to_write = months_data
    else:
        layout, category_resolver = _exporter_config(conn)
        to_write = process_monthly_data(
            months_data, category_resolver=category_resolver, sheet_categories=layout
        )
    for year, month in sorted(to_write.keys()):
        _write_month_sheet(wb, year, month, to_write[(year, month)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _month_transactions(conn, year, month):
    if month == 12:
        start, end = f"{year}-12-01", f"{year + 1}-01-01"
    else:
        start, end = f"{year}-{month:02d}-01", f"{year}-{month + 1:02d}-01"
    return get_transactions_by_date_range(conn, start, end)
